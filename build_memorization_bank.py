"""根据 analysis.ipynb 的背题规则筛选题库，保留原始列、值和顺序。

依赖：pip install openpyxl
运行：python build_memorization_bank.py --input "人工智能训练师.xlsx"
默认在脚本目录及其父目录查找输入，输出到脚本目录的 outputs/memorization。
多选：至少有一个错误选项，且每一个错误选项都含关键词，才剔除。
判断：题干含关键词且标准答案为错误，才剔除；字母答案按选项文本解析。
单选：默认仅剔除正确答案唯一最长的题；--single-rule original 可复现
原 notebook 的并列时取首项规则；--single-rule keep 可保留全部单选。
关键词沿用 notebook 的子串匹配口径（包括“仅”命中“不仅”）。
正确选项含关键词不阻止多选剔除，因为这里只检查所有错误选项。
生成一个工作簿：需要背的题库、已剔除（含原因）、筛选统计。
"""

from __future__ import annotations

import argparse
from collections import Counter
from pathlib import Path
import re
import unicodedata
from math import ceil

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ABSOLUTE_WORDS = [
    "仅", "无需", "一定", "始终", "绝对", "不能", "全部", "完全", "无限制",
    "任何", "不做任何", "任何情况下", "必然", "必须", "绝对不能",
    "没有任何", "完全不", "毫无", "随意", "永远",
]
FALSE_LABELS = {"错", "错误", "不正确", "否", "false", "f", "×", "✗", "✘"}
TRUE_LABELS = {"对", "正确", "是", "true", "t", "√", "✓", "✔"}


def text(value):
    return "" if value is None else str(value).strip()


def normalize(value):
    return unicodedata.normalize("NFKC", text(value))


def keyword_hits(value):
    return [word for word in ABSOLUTE_WORDS if word in text(value)]


def parse_answers(value):
    answer = re.sub(r"[\s,，、;；/|]+", "", normalize(value).upper())
    if not re.fullmatch(r"[A-H]+", answer):
        return None
    if len(answer) != len(set(answer)):
        return None
    return set(answer)


def resolve_columns(headers):
    def find(prefix):
        found = [i for i, h in enumerate(headers) if re.sub(r"\s+", "", text(h)).startswith(prefix)]
        if len(found) != 1:
            raise ValueError(f"必须存在唯一的{prefix}列，实际找到 {len(found)} 个")
        return found[0]

    options = {}
    for i, header in enumerate(headers):
        match = re.match(r"^选项([A-H])(?:$|[^A-Za-z])", re.sub(r"\s+", "", normalize(header)))
        if match:
            code = match[1]
            if code in options:
                raise ValueError(f"重复的选项列：{code}")
            options[code] = i
    if not options:
        raise ValueError("没有找到选项列")
    return find("题干"), find("题型"), find("正确答案"), dict(sorted(options.items()))


def classify(row, columns, single_rule="unique"):
    """返回 (是否剔除, 原因, 数据异常说明)。异常数据一律保留。"""
    stem_col, type_col, answer_col, option_cols = columns
    kind = text(row[type_col])
    options = {code: text(row[i]) for code, i in option_cols.items() if text(row[i])}
    raw_answer = row[answer_col]
    if not text(row[stem_col]):
        return False, "", "题干为空"

    if kind in {"判断", "判断题"}:
        value = normalize(raw_answer).lower()
        if value.upper() in option_cols:
            value = normalize(options.get(value.upper(), "")).lower()
        if value not in FALSE_LABELS | TRUE_LABELS:
            return False, "", "无法解析判断题答案"
        hits = keyword_hits(row[stem_col])
        if value in FALSE_LABELS and hits:
            return True, "判断答案为错误；题干命中：" + "、".join(hits), ""
        return False, "", ""

    if kind not in {"单选", "单选题", "多选", "多选题"}:
        return False, "", "未支持的题型，已保留"
    answers = parse_answers(raw_answer)
    if not answers or not answers <= options.keys():
        return False, "", "答案为空、格式异常或指向空选项"
    if len(options) < 2:
        return False, "", "有效选项少于两个"

    if kind in {"多选", "多选题"}:
        if len(answers) < 2:
            return False, "", "多选题正确答案少于两个"
        incorrect = [code for code in options if code not in answers]
        # 不允许 all([]) 把全选题误剔除。
        hits = {code: keyword_hits(options[code]) for code in incorrect}
        if incorrect and all(hits.values()):
            reason = "所有错误选项均命中：" + "；".join(
                f"{code}（{'、'.join(words)}）" for code, words in hits.items()
            )
            return True, reason, ""
        return False, "", ""

    if len(answers) != 1:
        return False, "", "单选题答案不是单个选项"
    if single_rule == "keep":
        return False, "", ""
    # original 精确沿用原 notebook 的 A-D 原始字符串长度口径。
    if single_rule == "original":
        if not all(code in options for code in "ABCD") or any(code in options for code in "EFGH"):
            return False, "", "原单选规则仅支持四个非空 A-D 选项"
        lengths = {code: len(str(row[option_cols[code]])) for code in "ABCD"}
    else:
        lengths = {code: len(value) for code, value in options.items()}
    longest = [code for code, length in lengths.items() if length == max(lengths.values())]
    hit = longest[0] in answers and (single_rule == "original" or len(longest) == 1)
    reason = "正确答案唯一最长" if single_rule == "unique" else "正确答案为首个最长选项（原规则）"
    return hit, reason if hit else "", ""


def style_sheet(sheet, headers):
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    widths = []
    for i, header in enumerate(headers, 1):
        label = text(header)
        width = 70 if "题干" in label else 38 if "选项" in label or "解析" in label else 22
        if "原因" in label or "说明" in label:
            width = 65
        widths.append(width)
        sheet.column_dimensions[get_column_letter(i)].width = width
    for row in sheet:
        line_count = 1
        for cell, width in zip(row, widths):
            cell.font = Font(name="Microsoft YaHei", size=10, bold=cell.row == 1,
                             color="FFFFFF" if cell.row == 1 else "222222")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="365778")
            value = text(cell.value)
            lines = sum(max(1, ceil(sum(2 if unicodedata.east_asian_width(c) in "WF" else 1 for c in line) / (width - 2))) for line in value.split("\n"))
            line_count = max(line_count, lines)
        sheet.row_dimensions[row[0].row].height = min(409, max(32, line_count * 16 + 8))


def run(input_path, output_path, single_rule="unique"):
    input_path, output_path = Path(input_path).resolve(), Path(output_path).resolve()
    if input_path == output_path:
        raise ValueError("输出路径不能覆盖原始题库")
    source = load_workbook(input_path, read_only=True, data_only=False)
    try:
        sheet = source.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        headers = list(next(rows))
        columns = resolve_columns(headers)
        kept, removed, anomalies = [], [], []
        totals, drops = Counter(), Counter()
        for number, row in enumerate(rows, 2):
            if not any(v is not None for v in row):
                continue
            row = list(row)
            kind = text(row[columns[1]])
            totals[kind] += 1
            drop, reason, anomaly = classify(row, columns, single_rule)
            if drop:
                removed.append(row + [number, reason])
                drops[kind] += 1
            else:
                kept.append(row)
            if anomaly:
                anomalies.append([number, kind, anomaly])
    finally:
        source.close()

    book = Workbook()
    main = book.active
    main.title = "需要背的题库"
    deleted = book.create_sheet("已剔除")
    summary = book.create_sheet("筛选统计")
    for target, labels, data in [
        (main, headers, kept),
        (deleted, headers + ["源Excel行号", "剔除原因"], removed),
    ]:
        target.append(labels)
        for row in data:
            target.append(row)
        style_sheet(target, labels)
    summary.append(["题型", "原始题数", "剔除题数", "需要背"])
    for kind, count in totals.items():
        summary.append([kind, count, drops[kind], count - drops[kind]])
    summary.append(["合计", sum(totals.values()), len(removed), len(kept)])
    summary.append([])
    summary.append(["规则", "说明"])
    summary.append(["单选", {"unique": "正确答案唯一最长才剔除；长度忽略首尾空白", "original": "原 notebook 规则：正确答案为 A-D 首个最长选项才剔除", "keep": "全部保留"}[single_rule]])
    summary.append(["多选", "至少存在一个错误选项，且所有错误选项均含关键词才剔除；全选题保留"])
    summary.append(["判断", "题干含关键词，且标准答案为错误才剔除；字母答案按选项文字解析"])
    summary.append(["匹配", "沿用 notebook 子串匹配；正确选项含关键词不影响多选规则"])
    summary.append(["关键词", "、".join(ABSOLUTE_WORDS)])
    summary.append(["异常题数", len(anomalies)])
    if anomalies:
        summary.append(["源Excel行号", "题型", "异常说明（均保留）"])
        for row in anomalies:
            summary.append(row)
    style_sheet(summary, ["题型", "说明", "剔除题数", "需要背"])
    output_path.parent.mkdir(parents=True, exist_ok=True)
    book.save(output_path)
    # 逐行核验导出结果，确认原始值、顺序和题量完整。
    check = load_workbook(output_path, read_only=True, data_only=False)
    try:
        assert list(check["需要背的题库"].values) == [tuple(headers)] + [tuple(r) for r in kept]
        assert check["已剔除"].max_row - 1 == len(removed)
        assert len(kept) + len(removed) == sum(totals.values())
    finally:
        check.close()
    for kind, count in totals.items():
        print(f"{kind}：原始 {count}，剔除 {drops[kind]}，保留 {count - drops[kind]}")
    print(f"合计：保留 {len(kept)}，剔除 {len(removed)}，异常保留 {len(anomalies)}")
    print(f"输出：{output_path}")
    return totals, drops


def main():
    base = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--input", type=Path, help="原始 xlsx，默认寻找人工智能训练师.xlsx")
    parser.add_argument("--output", type=Path, default=base / "outputs" / "memorization" / "人工智能训练师-完整背题库.xlsx")
    parser.add_argument("--single-rule", choices=["unique", "original", "keep"], default="unique")
    args = parser.parse_args()
    if args.input is None:
        candidates = [base / "人工智能训练师.xlsx", base.parent / "人工智能训练师.xlsx"]
        args.input = next((p for p in candidates if p.is_file()), None)
        if args.input is None:
            parser.error("没有找到原始题库，请用 --input 指定路径")
    run(args.input, args.output, args.single_rule)


if __name__ == "__main__":
    main()
